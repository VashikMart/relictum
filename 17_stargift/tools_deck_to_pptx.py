import json, os, re, shutil
from PIL import Image, PngImagePlugin
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE

import argparse
_ap=argparse.ArgumentParser(description='Собирает редактируемый PPTX по геометрии из tools_deck_extract.py')
_ap.add_argument('build_dir', help='папка, куда писал tools_deck_extract.py')
_ap.add_argument('--out', required=True, help='путь к .pptx')
_ap.add_argument('--root', default=os.path.dirname(os.path.abspath(__file__)), help='корень с картинками дека')
_A=_ap.parse_args()
SP=os.path.abspath(_A.build_dir)
ROOT=os.path.abspath(_A.root)
GEOM=json.load(open(os.path.join(SP,'geom.json'),encoding='utf-8'))
MEDIA=os.path.join(SP,'media'); shutil.rmtree(MEDIA,ignore_errors=True); os.makedirs(MEDIA)

PX=lambda v: Emu(int(round(v/96*914400)))
PT=lambda px: Pt(round(px*0.75,1))

def parse(css):
    m=re.match(r'rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?',css)
    if not m: return (0,0,0,1.0)
    a=float(m.group(4)) if m.group(4) is not None else 1.0
    return (int(float(m.group(1))),int(float(m.group(2))),int(float(m.group(3))),a)

def rgb(css):
    r,g,b,_=parse(css); return RGBColor(r,g,b)

def blend(css, bg_css):
    r,g,b,a=parse(css); br,bg_,bb,_=parse(bg_css)
    return RGBColor(int(round(r*a+br*(1-a))),int(round(g*a+bg_*(1-a))),int(round(b*a+bb*(1-a))))

ALIGN={'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT,'left':PP_ALIGN.LEFT,'start':PP_ALIGN.LEFT}

# ---------- white logo ----------
LOGO_W=os.path.join(SP,'logo_white.png')
lg=Image.open(os.path.join(ROOT,'img/stargift_logo.png')).convert('RGBA')
px=lg.load()
for y in range(lg.height):
    for x in range(lg.width):
        r,g,b,a=px[x,y]; px[x,y]=(255,255,255,a)
lg.save(LOGO_W)

seq=[0]
def uniq(src, box_px=None):
    """copy each placement to its own media file so no viewer can dedupe/reorder.

    Одного разного имени файла мало: python-pptx склеивает части по sha1
    содержимого, и две одинаковые картинки становятся одной media на весь дек —
    Keynote потом волен показать её не там, где надо. Поэтому в каждый файл
    кладём уникальную текстовую метку: байты разные, картинка та же."""
    seq[0]+=1
    tag='stargift-%03d'%seq[0]
    ext=os.path.splitext(src)[1].lower()
    dst=os.path.join(MEDIA,'m%03d%s'%(seq[0],ext if ext in ('.png','.jpg') else '.png'))
    im=Image.open(src)
    if im.mode in ('RGBA','LA','P') and ext=='.png':
        meta=PngImagePlugin.PngInfo(); meta.add_text('Placement', tag)
        im.convert('RGBA').save(dst, pnginfo=meta)
    else:
        im.convert('RGB').save(dst, quality=92, comment=tag.encode())
    return dst

def add_pic(slide, src, x,y,w,h, crop=None):
    p=slide.shapes.add_picture(uniq(src), PX(x),PX(y),PX(w),PX(h))
    if crop:
        p.crop_left,p.crop_right,p.crop_top,p.crop_bottom=crop
    return p

def fit_rect(nat, box, fit, pos):
    """returns (x,y,w,h, crop) in px within box (bx,by,bw,bh)"""
    bx,by,bw,bh=box; nw,nh=nat
    if nw==0 or nh==0: return (bx,by,bw,bh,None)
    px_,py_=0.5,0.5
    m=re.findall(r'([\d.]+)%',pos or '')
    if len(m)>=2: px_,py_=float(m[0])/100,float(m[1])/100
    if fit=='cover':
        s=max(bw/nw,bh/nh); dw,dh=nw*s,nh*s
        cl=(dw-bw)/dw*px_; cr=(dw-bw)/dw*(1-px_)
        ct=(dh-bh)/dh*py_; cb=(dh-bh)/dh*(1-py_)
        return (bx,by,bw,bh,(cl,cr,ct,cb))
    s=min(bw/nw,bh/nh); dw,dh=nw*s,nh*s
    return (bx+(bw-dw)/2, by+(bh-dh)/2, dw, dh, None)

def add_text(slide, t, override_runs=None, rect=None, lh=None):
    r=rect or t['rect']
    t=dict(t)
    if lh: t['lh']=lh
    box=slide.shapes.add_textbox(PX(r['x']-2),PX(r['y']-2),PX(r['w']+6),PX(r['h']+8))
    tf=box.text_frame
    tf.word_wrap=True
    tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    paras=[[]]
    for run in (override_runs or t['runs']):
        if run.get('br'): paras.append([]); continue
        paras[-1].append(run)
    first=True
    for pr in paras:
        # collapse source-code whitespace but keep NBSP (U+00A0) intact
        for i,run in enumerate(pr):
            run['text']=re.sub(r'[ \t\r\n\f\v]+',' ',run['text'])
        if pr:
            pr[0]['text']=pr[0]['text'].lstrip(' ')
            pr[-1]['text']=pr[-1]['text'].rstrip(' ')
        pr[:] = [r for r in pr if r['text']]
        if not pr: continue
        p=tf.paragraphs[0] if first else tf.add_paragraph()
        first=False
        p.alignment=ALIGN.get(t['align'],PP_ALIGN.LEFT)
        base=pr[0]['size'] if pr else t['base']
        if t['lh']: p.line_spacing=Pt(round(t['lh']*0.75,2))   # absolute: PPT multiples are font-metric based
        for run in pr:
            txt=run['text']
            if run.get('upper'): txt=txt.upper()
            rr=p.add_run(); rr.text=txt
            f=rr.font
            f.name=run['font']; f.size=PT(run['size'])
            f.bold = run['weight']>=600
            f.italic = run.get('italic',False)
            f.color.rgb=rgb(run['color'])
            if run.get('spacing'):
                f._rPr.set('spc', str(int(round(run['spacing']*0.75*100))))
            rPr=f._rPr
            for tag in ('latin','ea','cs'):
                el=rPr.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag,{'typeface':run['font']})
                rPr.append(el)
    return box

prs=Presentation()
prs.slide_width=PX(1280); prs.slide_height=PX(720)
blank=prs.slide_layouts[6]

for sid in sorted(GEOM.keys()):
    g=GEOM[sid]
    sl=prs.slides.add_slide(blank)
    # background
    bg=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,PX(1280),PX(720))
    bg.fill.solid(); bg.fill.fore_color.rgb=rgb(g['bg']); bg.line.fill.background()
    bg.shadow.inherit=False

    for d in g.get('decor',[]):
        r=d['rect']
        if r['w']>=1276 and r['h']>=716: continue          # the slide's own ground
        shape = MSO_SHAPE.OVAL if d.get('oval') else MSO_SHAPE.RECTANGLE
        sh=sl.shapes.add_shape(shape, PX(r['x']),PX(r['y']),
                               PX(max(r['w'],0.75)),PX(max(r['h'],0.75)))
        sh.fill.solid(); sh.fill.fore_color.rgb=blend(d['color'], g['bg'])
        sh.line.fill.background(); sh.shadow.inherit=False

    plates=g.get('plates') or ([g['plate']] if g.get('plate') else [])
    for plate in plates:
        pr=plate['rect']
        add_pic(sl, os.path.join(SP,'plates',plate['file']), pr['x'],pr['y'],pr['w'],pr['h'])

    def in_plate(r):
        for p in plates:
            q=p['rect']
            if (r['x']>=q['x']-1 and r['y']>=q['y']-1
                    and r['x']+r['w']<=q['x']+q['w']+1
                    and r['y']+r['h']<=q['y']+q['h']+1):
                return True
        return False

    for im in g['imgs']:
        r=im['rect']
        if in_plate(r): continue          # картинка уже впечатана в плейт
        src=os.path.join(ROOT, im['src'])
        if 'invert' in (im.get('filter') or ''): src=LOGO_W
        x,y,w,h,crop=fit_rect(im['nat'],(r['x'],r['y'],r['w'],r['h']),im['fit'],im['pos'])
        add_pic(sl,src,x,y,w,h,crop)

    for t in g['texts']:
        brs=sum(1 for r in t['runs'] if r.get('br'))
        if brs>=10:   # the two-column player list on s11
            names=[r for r in t['runs'] if not r.get('br')]
            half=(len(names)+1)//2
            r=t['rect']; colw=(r['w']-34)/2
            for k,part in enumerate((names[:half],names[half:])):
                runs=[]
                for i,n in enumerate(part):
                    if i: runs.append({'br':True})
                    runs.append(n)
                add_text(sl,t,runs,{'x':r['x']+k*(colw+34),'y':r['y'],'w':colw,'h':r['h']},
                         lh=r['h']/max(len(part),1))
            continue
        add_text(sl,t)

out=os.path.abspath(_A.out)
prs.save(out)
print('saved',out, os.path.getsize(out)//1024,'KB','slides',len(prs.slides.__iter__.__self__._sldIdLst))
